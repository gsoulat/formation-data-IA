"""Génère deux quiz Kahoot (format d'import officiel) pour le parcours Data Analyst :
- Kahoot_DataAnalyst_Maths_25Questions.xlsx : moyenne, médiane, quartiles, écart-type, % (cours 00 à 02)
- Kahoot_DataAnalyst_Git_Bash_25Questions.xlsx : commandes de base (semaine 1), niveau facile

Usage : uv run --with openpyxl python build_kahoot_da.py
Chaque question est écrite avec la BONNE réponse en premier ; le script mélange l'ordre (graine fixe)
et contrôle les limites Kahoot : question ≤ 120 caractères, réponse ≤ 75, durée autorisée.
"""
from pathlib import Path
import random
import shutil

import openpyxl

HERE = Path(__file__).parent
TEMPLATE = HERE / "KahootQuizTemplate.xlsx"
DUREES = {5, 10, 20, 30, 60, 90, 120, 240}

MATHS = [
    ("Quelle est la moyenne de 2, 4 et 6 ?", ["4", "3", "6", "12"], 20),
    ("Quelle est la médiane de 1, 3, 7, 8, 20 ?", ["7", "3", "8", "7,8"], 20),
    ("Quelle est la médiane de 2, 4, 6, 8 ?", ["5", "4", "6", "20"], 30),
    ("Quel est le mode de 3, 5, 5, 7, 9 ?", ["5", "3", "7", "9"], 20),
    ("Quel indicateur ne bouge presque pas si on ajoute une valeur énorme ?", ["La médiane", "La moyenne", "L'étendue", "L'écart-type"], 20),
    ("La moyenne est bien plus grande que la médiane. La distribution est…", ["Étalée à droite", "Symétrique", "Étalée à gauche", "Impossible à dire"], 30),
    ("Quelle est l'étendue de 4, 9, 15, 30 ?", ["26", "30", "15", "34"], 20),
    ("Le 1er quartile (Q1) laisse en dessous de lui…", ["25 % des valeurs", "10 % des valeurs", "50 % des valeurs", "75 % des valeurs"], 20),
    ("La médiane est aussi appelée…", ["Le 2e quartile", "Le 1er quartile", "Le 3e quartile", "Le mode"], 20),
    ("Entre Q1 et Q3, on trouve…", ["La moitié des valeurs", "Toutes les valeurs", "Un quart des valeurs", "Les valeurs extrêmes"], 20),
    ("Que mesure l'écart-type ?", ["La dispersion autour de la moyenne", "Le centre des données", "La valeur la plus fréquente", "Le nombre de valeurs"], 20),
    ("Pour la variance, pourquoi met-on les écarts au carré ?", ["Pour que les + et les − ne s'annulent pas", "Pour aller plus vite", "Pour arrondir", "Pour trier les valeurs"], 30),
    ("L'écart-type est la racine carrée de…", ["La variance", "La moyenne", "La médiane", "L'étendue"], 20),
    ("Série 2, 4, 4, 4, 5, 5, 7, 9 (moyenne 5). Écart-type (divisé par n) ?", ["2", "1", "4", "5"], 60),
    ("Quelle fonction Google Sheets calcule la médiane ?", ["MEDIAN", "AVERAGE", "MODE", "COUNT"], 20),
    ("Que fait la fonction COUNTBLANK (NB.VIDE) ?", ["Elle compte les cellules vides", "Elle compte les nombres", "Elle additionne", "Elle compte les textes"], 20),
    ("La fonction AVERAGE (MOYENNE) ignore les cellules vides.", ["Vrai", "Faux"], 10),
    ("Un prix passe de 50 € à 60 €. Quelle est l'évolution ?", ["+20 %", "+10 %", "+16,7 %", "+60 %"], 30),
    ("Un taux passe de 16 % à 19 %. Il augmente de…", ["3 points", "3 %", "19 %", "16 points"], 30),
    ("30 commandes livrées sur 120. Quelle part est livrée ?", ["25 %", "30 %", "40 %", "12 %"], 20),
    ("Moyenne pondérée de 10 (coefficient 1) et 16 (coefficient 2) ?", ["14", "13", "12", "16"], 60),
    ("Pour comparer des territoires, on pondère souvent par…", ["La population", "La superficie", "L'ordre alphabétique", "Le code postal"], 20),
    ("Quel graphique montre la forme d'une distribution ?", ["L'histogramme", "Le camembert", "La courbe", "La carte"], 20),
    ("Moyenne 1 680 €, médiane 177 €. Pour décrire une commande type, on donne…", ["La médiane", "La moyenne", "Le maximum", "La somme"], 20),
    ("« Note moyenne de nos 613 commandes : 4,1 » alors que 560 sont notées. Le souci ?", ["Le périmètre est faux", "Le calcul est faux", "4,1 est trop haut", "Aucun"], 30),
]

GIT_BASH = [
    ("Quelle commande affiche le dossier dans lequel tu te trouves ?", ["pwd", "ls", "cd", "mkdir"], 20),
    ("Quelle commande liste les fichiers d'un dossier ?", ["ls", "cd", "rm", "cat"], 20),
    ("Que fait la commande « cd .. » ?", ["Elle remonte d'un dossier", "Elle supprime un dossier", "Elle crée un dossier", "Elle liste les fichiers"], 20),
    ("Comment créer un dossier nommé data ?", ["mkdir data", "touch data", "cd data", "rm data"], 20),
    ("Comment créer un fichier vide notes.txt ?", ["touch notes.txt", "mkdir notes.txt", "ls notes.txt", "pwd notes.txt"], 20),
    ("Quelle commande affiche le contenu d'un fichier texte ?", ["cat", "cd", "mkdir", "pwd"], 20),
    ("Que fait « rm fichier.txt » ?", ["Il supprime le fichier", "Il renomme le fichier", "Il copie le fichier", "Il ouvre le fichier"], 20),
    ("Comment copier a.txt vers b.txt ?", ["cp a.txt b.txt", "mv a.txt b.txt", "rm a.txt b.txt", "cat a.txt b.txt"], 20),
    ("Quelle commande renomme ou déplace un fichier ?", ["mv", "cp", "ls", "rm"], 20),
    ("Dans le terminal, que désigne le symbole ~ ?", ["Ton dossier personnel", "La racine du disque", "Le dossier parent", "Un fichier caché"], 20),
    ("Quelle touche complète automatiquement un nom de fichier ?", ["Tab", "Entrée", "Échap", "Espace"], 10),
    ("Git, c'est…", ["Un logiciel de gestion de versions", "Un langage de programmation", "Un réseau social", "Un tableur"], 20),
    ("GitHub, c'est…", ["Un site qui héberge des dépôts Git", "Le même logiciel que Git", "Un éditeur de texte", "Un navigateur web"], 20),
    ("Quelle commande crée un dépôt Git dans le dossier courant ?", ["git init", "git start", "git new", "git create"], 20),
    ("Quelle commande montre les fichiers modifiés ?", ["git status", "git log", "git push", "git clone"], 20),
    ("Quelle commande prépare un fichier pour le prochain commit ?", ["git add", "git commit", "git push", "git pull"], 20),
    ("Comment enregistrer une version avec un message ?", ['git commit -m "message"', 'git save "message"', 'git add -m "message"', 'git push "message"'], 30),
    ("Quelle commande envoie tes commits sur GitHub ?", ["git push", "git pull", "git init", "git status"], 20),
    ("Quelle commande récupère les nouveautés depuis GitHub ?", ["git pull", "git push", "git add", "git commit"], 20),
    ("Comment copier un dépôt GitHub sur ton ordinateur ?", ["git clone", "git copy", "git download", "git init"], 20),
    ("Quelle commande affiche l'historique des commits ?", ["git log", "git status", "git history", "git show-all"], 20),
    ("Un bon message de commit…", ["Dit ce que fait le commit", "Tient en un mot : « modifs »", "Est vide", "Contient seulement la date"], 20),
    ("À quoi sert le fichier .gitignore ?", ["Lister les fichiers à ne pas versionner", "Cacher le dépôt", "Supprimer des fichiers", "Écrire la documentation"], 30),
    ("Que présente le fichier README.md d'un dépôt ?", ["Le projet et comment l'utiliser", "Les mots de passe", "L'historique des commits", "La liste des branches"], 20),
    ("À quoi sert une branche Git ?", ["Travailler sur une version parallèle", "Supprimer le dépôt", "Envoyer un e-mail", "Trier les fichiers"], 20),
]


def construire(questions, nom, graine):
    assert len(questions) == 25, (nom, len(questions))
    rnd = random.Random(graine)
    dest = HERE / nom
    shutil.copy(TEMPLATE, dest)
    wb = openpyxl.load_workbook(dest)
    ws = wb.active
    positions = []
    for i, (q, reps, duree) in enumerate(questions):
        assert len(q) <= 120, (q, len(q))
        assert all(len(r) <= 75 for r in reps), reps
        assert duree in DUREES
        bonne = reps[0]
        ordre = reps[:]
        if len(ordre) == 4:
            rnd.shuffle(ordre)
        r = 9 + i
        ws.cell(r, 1, i + 1)
        ws.cell(r, 2, q)
        for j in range(4):
            ws.cell(r, 3 + j, ordre[j] if j < len(ordre) else None)
        ws.cell(r, 7, duree)
        ws.cell(r, 8, str(ordre.index(bonne) + 1))
        positions.append(ordre.index(bonne) + 1)
    wb.save(dest)
    print(f"{nom} : 25 questions · bonnes réponses en position 1/2/3/4 : "
          f"{[positions.count(k) for k in (1, 2, 3, 4)]}")


construire(MATHS, "Kahoot_DataAnalyst_Maths_25Questions.xlsx", 11)
construire(GIT_BASH, "Kahoot_DataAnalyst_Git_Bash_25Questions.xlsx", 22)
